import os
import re
import uuid
import subprocess
import shutil
from pathlib import Path

from flask import Flask, request, send_file, jsonify, render_template, Response
from flask_cors import CORS

from pdf2docx import Converter as PdfToDocxConverter
from docx import Document
from openpyxl import Workbook
from PIL import Image
import pytesseract

from seo_content import TOOLS, COMPRESS_SIZE_PAGES, SITE_URL

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
OUTPUT_DIR = BASE_DIR / "outputs"
UPLOAD_DIR.mkdir(exist_ok=True)
OUTPUT_DIR.mkdir(exist_ok=True)

MAX_FILE_SIZE_MB = 25

app = Flask(__name__)
CORS(app)
app.config["MAX_CONTENT_LENGTH"] = MAX_FILE_SIZE_MB * 1024 * 1024


# ---------------------------------------------------------------------------
# Helper utilities (यूनिकोड और क्लीनअप बग फिक्स)
# ---------------------------------------------------------------------------

def save_upload(file_storage, allowed_ext):
    """
    secure_filename हिंदी/यूनिकोड अक्षरों को पूरी तरह मिटा देता है,
    इसलिए एक्सटेंशन मूल नाम से निकालकर सुरक्षित UUID पाथ पर सेव किया गया है।
    """
    orig_name = file_storage.filename or ""
    ext = orig_name.rsplit(".", 1)[-1].lower() if "." in orig_name else ""
    if ext not in allowed_ext:
        raise ValueError(f"सिर्फ {', '.join(allowed_ext)} फ़ाइलें ही मान्य हैं।")
    
    job_id = uuid.uuid4().hex
    saved_path = UPLOAD_DIR / f"{job_id}.{ext}"
    file_storage.save(saved_path)
    return saved_path, job_id


def cleanup(*paths):
    """फ़ाइलों और फ़ोल्डरों को सुरक्षित रूप से साफ़ करता है।"""
    for p in paths:
        try:
            if p:
                target = Path(p)
                if target.is_dir():
                    shutil.rmtree(target, ignore_errors=True)
                elif target.exists():
                    target.unlink()
        except Exception:
            pass


# ---------------------------------------------------------------------------
# 1) PDF -> Word (Auto-cleanup जोड़ा गया)
# ---------------------------------------------------------------------------

@app.route("/api/pdf-to-word", methods=["POST"])
def pdf_to_word():
    if "file" not in request.files:
        return jsonify({"error": "कोई फ़ाइल नहीं भेजी गई।"}), 400

    try:
        src_path, job_id = save_upload(request.files["file"], {"pdf"})
    except ValueError as e:
        return jsonify({"error": str(e)}), 400

    out_path = OUTPUT_DIR / f"{job_id}.docx"
    try:
        converter = PdfToDocxConverter(str(src_path))
        converter.convert(str(out_path))
        converter.close()
    except Exception as e:
        cleanup(src_path, out_path)
        return jsonify({"error": f"Conversion असफल: {e}"}), 500
    finally:
        cleanup(src_path)

    response = send_file(
        out_path,
        as_attachment=True,
        download_name="converted.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    # मेमोरी/डिस्क लीक रोकने के लिए डाउनलोड के बाद आउटपुट फ़ाइल डिलीट करना
    response.call_on_close(lambda: cleanup(out_path))
    return response


# ---------------------------------------------------------------------------
# 2) Word -> PDF (कॉन्करेंट इंस्टॉलेशन लॉक फिक्स)
# ---------------------------------------------------------------------------

@app.route("/api/word-to-pdf", methods=["POST"])
def word_to_pdf():
    if "file" not in request.files:
        return jsonify({"error": "कोई फ़ाइल नहीं भेजी गई।"}), 400

    try:
        src_path, job_id = save_upload(request.files["file"], {"doc", "docx"})
    except ValueError as e:
        return jsonify({"error": str(e)}), 400

    job_out_dir = OUTPUT_DIR / job_id
    job_out_dir.mkdir(exist_ok=True)
    user_inst_dir = UPLOAD_DIR / f"soffice_user_{job_id}"

    try:
        result = subprocess.run(
            [
                "soffice",
                f"-env:UserInstallation=file://{user_inst_dir}",
                "--headless", "--norestore",
                "--convert-to", "pdf",
                "--outdir", str(job_out_dir),
                str(src_path),
            ],
            capture_output=True, text=True, timeout=90,
        )
        produced = list(job_out_dir.glob("*.pdf"))
        if result.returncode != 0 or not produced:
            raise RuntimeError(result.stderr or "LibreOffice ने PDF नहीं बनाई।")
        out_path = produced[0]
    except FileNotFoundError:
        cleanup(src_path, job_out_dir, user_inst_dir)
        return jsonify({"error": "Server पर LibreOffice (soffice) मौजूद नहीं है।"}), 500
    except Exception as e:
        cleanup(src_path, job_out_dir, user_inst_dir)
        return jsonify({"error": f"Conversion असफल: {e}"}), 500
    finally:
        cleanup(src_path, user_inst_dir)

    response = send_file(
        out_path, as_attachment=True,
        download_name="converted.pdf", mimetype="application/pdf"
    )
    response.call_on_close(lambda: cleanup(job_out_dir))
    return response


# ---------------------------------------------------------------------------
# 3) PDF Compress (ऑटो-क्लीनअप और रिसोर्स मैनेजमेंट)
# ---------------------------------------------------------------------------

GS_QUALITY_PRESETS = {
    "low": "/screen",
    "medium": "/ebook",
    "high": "/printer",
}

GS_TARGET_ATTEMPTS = [
    ("/printer", 200),
    ("/ebook", 150),
    ("/ebook", 120),
    ("/screen", 96),
    ("/screen", 72),
    ("/screen", 50),
    ("/screen", 36),
]


def run_ghostscript(src_path, out_path, preset, resolution=None):
    cmd = [
        "gs", "-sDEVICE=pdfwrite", "-dCompatibilityLevel=1.4",
        f"-dPDFSETTINGS={preset}",
        "-dNOPAUSE", "-dQUIET", "-dBATCH",
        "-dDetectDuplicateImages=true", "-dCompressFonts=true",
    ]
    if resolution:
        cmd += [
            "-dDownsampleColorImages=true", f"-dColorImageResolution={resolution}",
            "-dDownsampleGrayImages=true", f"-dGrayImageResolution={resolution}",
            "-dDownsampleMonoImages=true", f"-dMonoImageResolution={resolution}",
        ]
    cmd += [f"-sOutputFile={out_path}", str(src_path)]
    result = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
    if result.returncode != 0 or not Path(out_path).exists():
        raise RuntimeError(result.stderr or "Ghostscript निष्पादन असफल रहा।")


@app.route("/api/pdf-compress", methods=["POST"])
def pdf_compress():
    if "file" not in request.files:
        return jsonify({"error": "कोई फ़ाइल नहीं भेजी गई।"}), 400

    quality = request.form.get("quality", "medium")
    target_kb_raw = request.form.get("target_kb", "").strip()
    target_kb = int(target_kb_raw) if target_kb_raw.isdigit() else None

    try:
        src_path, job_id = save_upload(request.files["file"], {"pdf"})
    except ValueError as e:
        return jsonify({"error": str(e)}), 400

    original_size = src_path.stat().st_size
    out_path = OUTPUT_DIR / f"{job_id}_compressed.pdf"
    reached_target = None
    temp_files = []

    try:
        if target_kb:
            target_bytes = target_kb * 1024
            best_path, best_size = None, None
            for i, (preset, res) in enumerate(GS_TARGET_ATTEMPTS):
                trial_path = OUTPUT_DIR / f"{job_id}_try{i}.pdf"
                temp_files.append(trial_path)
                run_ghostscript(src_path, trial_path, preset, res)
                trial_size = trial_path.stat().st_size
                if best_path is None or trial_size < best_size:
                    best_path, best_size = trial_path, trial_size
                if trial_size <= target_bytes:
                    break
            shutil.copy(str(best_path), str(out_path))
            reached_target = best_size <= target_bytes
        else:
            gs_setting = GS_QUALITY_PRESETS.get(quality, "/ebook")
            run_ghostscript(src_path, out_path, gs_setting)
    except FileNotFoundError:
        cleanup(src_path, out_path, *temp_files)
        return jsonify({"error": "Server पर Ghostscript (gs) मौजूद नहीं है।"}), 500
    except Exception as e:
        cleanup(src_path, out_path, *temp_files)
        return jsonify({"error": f"Compression असफल: {e}"}), 500
    finally:
        cleanup(src_path, *temp_files)

    compressed_size = out_path.stat().st_size
    if compressed_size >= original_size and src_path.exists():
        shutil.copy(src_path, out_path)
        compressed_size = original_size

    saved_percent = round((1 - compressed_size / original_size) * 100) if original_size else 0

    response = send_file(
        out_path, as_attachment=True,
        download_name="compressed.pdf", mimetype="application/pdf"
    )
    response.headers["X-Original-Size"] = str(original_size)
    response.headers["X-Compressed-Size"] = str(compressed_size)
    response.headers["X-Saved-Percent"] = str(saved_percent)
    if target_kb:
        response.headers["X-Target-Reached"] = "true" if reached_target else "false"
    response.headers["Access-Control-Expose-Headers"] = (
        "X-Original-Size, X-Compressed-Size, X-Saved-Percent, X-Target-Reached"
    )
    response.call_on_close(lambda: cleanup(out_path))
    return response


# ---------------------------------------------------------------------------
# 4) Image -> Word (Auto-cleanup जोड़ा गया)
# ---------------------------------------------------------------------------

@app.route("/api/image-to-word", methods=["POST"])
def image_to_word():
    if "file" not in request.files:
        return jsonify({"error": "कोई फ़ाइल नहीं भेजी गई।"}), 400

    try:
        src_path, job_id = save_upload(
            request.files["file"], {"png", "jpg", "jpeg", "webp", "bmp"}
        )
    except ValueError as e:
        return jsonify({"error": str(e)}), 400

    out_path = OUTPUT_DIR / f"{job_id}.docx"
    try:
        img = Image.open(src_path)
        text = pytesseract.image_to_string(img, lang="eng+hin")
        doc = Document()
        doc.add_heading("Extracted Text", level=1)
        for line in text.split("\n"):
            if line.strip():
                doc.add_paragraph(line)
        doc.save(out_path)
    except pytesseract.TesseractNotFoundError:
        cleanup(src_path, out_path)
        return jsonify({"error": "Server पर Tesseract OCR मौजूद नहीं है।"}), 500
    except Exception as e:
        cleanup(src_path, out_path)
        return jsonify({"error": f"Conversion असफल: {e}"}), 500
    finally:
        cleanup(src_path)

    response = send_file(
        out_path, as_attachment=True, download_name="extracted.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    response.call_on_close(lambda: cleanup(out_path))
    return response


# ---------------------------------------------------------------------------
# 5) Image -> Excel (Auto-cleanup जोड़ा गया)
# ---------------------------------------------------------------------------

def extract_table_rows(img):
    data = pytesseract.image_to_data(img, lang="eng+hin", output_type=pytesseract.Output.DICT)
    lines = {}
    for i in range(len(data["text"])):
        text = data["text"][i].strip()
        if not text:
            continue
        key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
        lines.setdefault(key, []).append(
            {"text": text, "left": data["left"][i], "width": data["width"][i]}
        )

    rows = []
    for key in sorted(lines.keys()):
        words = sorted(lines[key], key=lambda w: w["left"])
        columns, current = [], [words[0]["text"]]
        prev = words[0]
        for w in words[1:]:
            gap = w["left"] - (prev["left"] + prev["width"])
            avg_char_width = (
                (prev["width"] / max(len(prev["text"]), 1))
                + (w["width"] / max(len(w["text"]), 1))
            ) / 2
            if gap > max(8, 2.2 * avg_char_width):
                columns.append(" ".join(current))
                current = [w["text"]]
            else:
                current.append(w["text"])
            prev = w
        columns.append(" ".join(current))
        rows.append(columns)
    return rows


@app.route("/api/image-to-excel", methods=["POST"])
def image_to_excel():
    if "file" not in request.files:
        return jsonify({"error": "कोई फ़ाइल नहीं भेजी गई।"}), 400

    try:
        src_path, job_id = save_upload(
            request.files["file"], {"png", "jpg", "jpeg", "webp", "bmp"}
        )
    except ValueError as e:
        return jsonify({"error": str(e)}), 400

    out_path = OUTPUT_DIR / f"{job_id}.xlsx"
    try:
        img = Image.open(src_path)
        rows = extract_table_rows(img)

        wb = Workbook()
        ws = wb.active
        ws.title = "Extracted Data"

        for row_idx, columns in enumerate(rows, start=1):
            for col_idx, value in enumerate(columns, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        wb.save(out_path)
    except pytesseract.TesseractNotFoundError:
        cleanup(src_path, out_path)
        return jsonify({"error": "Server पर Tesseract OCR मौजूद नहीं है।"}), 500
    except Exception as e:
        cleanup(src_path, out_path)
        return jsonify({"error": f"Conversion असफल: {e}"}), 500
    finally:
        cleanup(src_path)

    response = send_file(
        out_path, as_attachment=True, download_name="extracted.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response.call_on_close(lambda: cleanup(out_path))
    return response


# ---------------------------------------------------------------------------
# Pages & Sitemap Routes
# ---------------------------------------------------------------------------

@app.route("/")
def home():
    return render_template("index.html", tools=TOOLS, site_url=SITE_URL)


@app.route("/<slug>")
def tool_page(slug):
    if slug in TOOLS:
        return render_template(
            "tool_page.html", slug=slug, tool=TOOLS[slug], all_tools=TOOLS,
            target_kb=None, site_url=SITE_URL,
        )
    if slug in COMPRESS_SIZE_PAGES:
        page = COMPRESS_SIZE_PAGES[slug]
        base_tool = dict(TOOLS["pdf-compress"])
        base_tool.update(page)
        return render_template(
            "tool_page.html", slug=slug, tool=base_tool, all_tools=TOOLS,
            target_kb=page["target_kb"], site_url=SITE_URL,
        )
    from flask import abort
    abort(404)


@app.route("/sitemap.xml")
def sitemap():
    urls = [SITE_URL + "/"]
    urls += [f"{SITE_URL}/{slug}" for slug in TOOLS]
    urls += [f"{SITE_URL}/{slug}" for slug in COMPRESS_SIZE_PAGES]
    xml_items = "".join(f"<url><loc>{u}</loc></url>" for u in urls)
    xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">'
        f"{xml_items}</urlset>"
    )
    return Response(xml, mimetype="application/xml")


@app.route("/robots.txt")
def robots():
    lines = ["User-agent: *", "Allow: /", f"Sitemap: {SITE_URL}/sitemap.xml"]
    return Response("\n".join(lines), mimetype="text/plain")


@app.route("/api/health")
def health():
    return jsonify({"status": "ok"})


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=False)

