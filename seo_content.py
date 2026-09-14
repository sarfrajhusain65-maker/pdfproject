"""
FileTools Backend
==================
असली working conversion tools:
1. PDF -> Word     (pdf2docx लाइब्रेरी)
2. Word -> PDF     (LibreOffice headless की मदद से)
3. Image -> Word   (Tesseract OCR + python-docx)
4. Image -> Excel  (Tesseract OCR + openpyxl, table की तरह arrange)
5. PDF Compress    (Ghostscript की मदद से size घटाना)

चलाने से पहले जरूरी system packages (Ubuntu/Debian VPS पर):
    sudo apt update
    sudo apt install -y libreoffice tesseract-ocr poppler-utils ghostscript python3-pip

Python packages: requirements.txt देखें
"""

import os
import re
import uuid
import subprocess
import shutil
from pathlib import Path

from flask import Flask, request, send_file, jsonify, render_template, Response, url_for
from flask_cors import CORS
from werkzeug.utils import secure_filename

from pdf2docx import Converter as PdfToDocxConverter
from docx import Document
from openpyxl import Workbook
from PIL import Image
import pytesseract

from seo_content import TOOLS, COMPRESS_SIZE_PAGES, SITE_URL

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
OUTPUT_DIR = BASE_DIR / "outputs"
UPLOAD_DIR.mkdir(exist_ok=True)
OUTPUT_DIR.mkdir(exist_ok=True)

MAX_FILE_SIZE_MB = 25

app = Flask(__name__)
CORS(app)
app.config["MAX_CONTENT_LENGTH"] = MAX_FILE_SIZE_MB * 1024 * 1024


# ---------------------------------------------------------------------------
# Helper utilities
# ---------------------------------------------------------------------------

def save_upload(file_storage, allowed_ext):
    filename = secure_filename(file_storage.filename)
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in allowed_ext:
        raise ValueError(f"सिर्फ {', '.join(allowed_ext)} फ़ाइलें ही allowed हैं।")
    job_id = uuid.uuid4().hex
    saved_path = UPLOAD_DIR / f"{job_id}.{ext}"
    file_storage.save(saved_path)
    return saved_path, job_id


def cleanup(*paths):
    for p in paths:
        try:
            if p and Path(p).exists():
                Path(p).unlink()
        except Exception:
            pass


# ---------------------------------------------------------------------------
# 1) PDF -> Word
# ---------------------------------------------------------------------------

@app.route("/api/pdf-to-word", methods=["POST"])
def pdf_to_word():
    if "file" not in request.files:
        return jsonify({"error": "कोई फ़ाइल नहीं भेजी गई।"}), 400

    try:
        src_path, job_id = save_upload(request.files["file"], {"pdf"})
    except ValueError as e:
        return jsonify({"error": str(e)}), 400

    out_path = OUTPUT_DIR / f"{job_id}.docx"
    try:
        converter = PdfToDocxConverter(str(src_path))
        converter.convert(str(out_path))
        converter.close()
    except Exception as e:
        cleanup(src_path)
        return jsonify({"error": f"Conversion असफल: {e}"}), 500

    cleanup(src_path)
    return send_file(
        out_path,
        as_attachment=True,
        download_name="converted.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


# ---------------------------------------------------------------------------
# 2) Word -> PDF  (LibreOffice headless चाहिए server पर)
# ---------------------------------------------------------------------------

@app.route("/api/word-to-pdf", methods=["POST"])
def word_to_pdf():
    if "file" not in request.files:
        return jsonify({"error": "कोई फ़ाइल नहीं भेजी गई।"}), 400

    try:
        src_path, job_id = save_upload(request.files["file"], {"doc", "docx"})
    except ValueError as e:
        return jsonify({"error": str(e)}), 400

    job_out_dir = OUTPUT_DIR / job_id
    job_out_dir.mkdir(exist_ok=True)

    try:
        result = subprocess.run(
            [
                "soffice", "--headless", "--norestore",
                "--convert-to", "pdf",
                "--outdir", str(job_out_dir),
                str(src_path),
            ],
            capture_output=True, text=True, timeout=90,
        )
        produced = list(job_out_dir.glob("*.pdf"))
        if result.returncode != 0 or not produced:
            raise RuntimeError(result.stderr or "LibreOffice ने PDF नहीं बनाई।")
        out_path = produced[0]
    except FileNotFoundError:
        cleanup(src_path)
        shutil.rmtree(job_out_dir, ignore_errors=True)
        return jsonify({
            "error": "Server पर LibreOffice (soffice) install नहीं है। README देखें।"
        }), 500
    except Exception as e:
        cleanup(src_path)
        shutil.rmtree(job_out_dir, ignore_errors=True)
        return jsonify({"error": f"Conversion असफल: {e}"}), 500

    cleanup(src_path)
    response = send_file(out_path, as_attachment=True,
                          download_name="converted.pdf", mimetype="application/pdf")
    response.call_on_close(lambda: shutil.rmtree(job_out_dir, ignore_errors=True))
    return response


# ---------------------------------------------------------------------------
# 3) PDF Compress (Ghostscript)
# ---------------------------------------------------------------------------

GS_QUALITY_PRESETS = {
    "low": "/screen",     # सबसे ज़्यादा compression, quality सबसे कम (web/screen)
    "medium": "/ebook",   # अच्छा balance (default)
    "high": "/printer",   # कम compression, quality बेहतर
}

# "Compress to X KB" के लिए — mild से aggressive तक कोशिशें,
# जब तक target size न मिल जाए या attempts खत्म न हो जाएं
GS_TARGET_ATTEMPTS = [
    ("/printer", 200),
    ("/ebook", 150),
    ("/ebook", 120),
    ("/screen", 96),
    ("/screen", 72),
    ("/screen", 50),
    ("/screen", 36),
]


def run_ghostscript(src_path, out_path, preset, resolution=None):
    cmd = [
        "gs", "-sDEVICE=pdfwrite", "-dCompatibilityLevel=1.4",
        f"-dPDFSETTINGS={preset}",
        "-dNOPAUSE", "-dQUIET", "-dBATCH",
        "-dDetectDuplicateImages=true", "-dCompressFonts=true",
    ]
    if resolution:
        cmd += [
            "-dDownsampleColorImages=true", f"-dColorImageResolution={resolution}",
            "-dDownsampleGrayImages=true", f"-dGrayImageResolution={resolution}",
            "-dDownsampleMonoImages=true", f"-dMonoImageResolution={resolution}",
        ]
    cmd += [f"-sOutputFile={out_path}", str(src_path)]
    result = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
    if result.returncode != 0 or not Path(out_path).exists():
        raise RuntimeError(result.stderr or "Ghostscript ने PDF नहीं बनाई।")


@app.route("/api/pdf-compress", methods=["POST"])
def pdf_compress():
    if "file" not in request.files:
        return jsonify({"error": "कोई फ़ाइल नहीं भेजी गई।"}), 400

    quality = request.form.get("quality", "medium")
    target_kb_raw = request.form.get("target_kb", "").strip()
    target_kb = int(target_kb_raw) if target_kb_raw.isdigit() else None

    try:
        src_path, job_id = save_upload(request.files["file"], {"pdf"})
    except ValueError as e:
        return jsonify({"error": str(e)}), 400

    original_size = src_path.stat().st_size
    out_path = OUTPUT_DIR / f"{job_id}_compressed.pdf"
    reached_target = None

    try:
        if target_kb:
            # हर attempt को अलग temp file में बनाकर सबसे अच्छा result चुनते हैं
            target_bytes = target_kb * 1024
            best_path, best_size = None, None
            for i, (preset, res) in enumerate(GS_TARGET_ATTEMPTS):
                trial_path = OUTPUT_DIR / f"{job_id}_try{i}.pdf"
                run_ghostscript(src_path, trial_path, preset, res)
                trial_size = trial_path.stat().st_size
                if best_path is None or trial_size < best_size:
                    if best_path:
                        cleanup(best_path)
                    best_path, best_size = trial_path, trial_size
                else:
                    cleanup(trial_path)
                if trial_size <= target_bytes:
                    break
            shutil.move(str(best_path), out_path)
            reached_target = best_size <= target_bytes
        else:
            gs_setting = GS_QUALITY_PRESETS.get(quality, "/ebook")
            run_ghostscript(src_path, out_path, gs_setting)
    except FileNotFoundError:
        cleanup(src_path)
        return jsonify({
            "error": "Server पर Ghostscript (gs) install नहीं है। README देखें।"
        }), 500
    except Exception as e:
        cleanup(src_path)
        return jsonify({"error": f"Compression असफल: {e}"}), 500

    compressed_size = out_path.stat().st_size
    # अगर compressed file बड़ी निकले (कभी-कभी already-compressed PDFs पर होता है)
    # तो original ही भेज दो ताकि user को नुकसान न हो
    if compressed_size >= original_size:
        shutil.copy(src_path, out_path)
        compressed_size = original_size

    cleanup(src_path)
    saved_percent = round((1 - compressed_size / original_size) * 100) if original_size else 0

    response = send_file(out_path, as_attachment=True,
                          download_name="compressed.pdf", mimetype="application/pdf")
    response.headers["X-Original-Size"] = str(original_size)
    response.headers["X-Compressed-Size"] = str(compressed_size)
    response.headers["X-Saved-Percent"] = str(saved_percent)
    if target_kb:
        response.headers["X-Target-Reached"] = "true" if reached_target else "false"
    response.headers["Access-Control-Expose-Headers"] = (
        "X-Original-Size, X-Compressed-Size, X-Saved-Percent, X-Target-Reached"
    )
    return response


# ---------------------------------------------------------------------------
# 4) Image -> Word (OCR)
# ---------------------------------------------------------------------------

@app.route("/api/image-to-word", methods=["POST"])
def image_to_word():
    if "file" not in request.files:
        return jsonify({"error": "कोई फ़ाइल नहीं भेजी गई।"}), 400

    try:
        src_path, job_id = save_upload(
            request.files["file"], {"png", "jpg", "jpeg", "webp", "bmp"}
        )
    except ValueError as e:
        return jsonify({"error": str(e)}), 400

    out_path = OUTPUT_DIR / f"{job_id}.docx"
    try:
        img = Image.open(src_path)
        text = pytesseract.image_to_string(img, lang="eng+hin")
        doc = Document()
        doc.add_heading("Extracted Text", level=1)
        for line in text.split("\n"):
            doc.add_paragraph(line)
        doc.save(out_path)
    except pytesseract.TesseractNotFoundError:
        cleanup(src_path)
        return jsonify({
            "error": "Server पर Tesseract OCR install नहीं है। README देखें।"
        }), 500
    except Exception as e:
        cleanup(src_path)
        return jsonify({"error": f"Conversion असफल: {e}"}), 500

    cleanup(src_path)
    return send_file(
        out_path, as_attachment=True, download_name="extracted.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


# ---------------------------------------------------------------------------
# 5) Image -> Excel (OCR + best-effort column detection)
# ---------------------------------------------------------------------------

def split_into_columns(line: str):
    """2 या ज्यादा spaces / tab को column-separator मानकर split करता है."""
    parts = re.split(r"\t+|\s{2,}", line.strip())
    return [p for p in parts if p != ""]


@app.route("/api/image-to-excel", methods=["POST"])
def image_to_excel():
    if "file" not in request.files:
        return jsonify({"error": "कोई फ़ाइल नहीं भेजी गई।"}), 400

    try:
        src_path, job_id = save_upload(
            request.files["file"], {"png", "jpg", "jpeg", "webp", "bmp"}
        )
    except ValueError as e:
        return jsonify({"error": str(e)}), 400

    out_path = OUTPUT_DIR / f"{job_id}.xlsx"
    try:
        img = Image.open(src_path)
        text = pytesseract.image_to_string(img, lang="eng+hin")

        wb = Workbook()
        ws = wb.active
        ws.title = "Extracted Data"

        row_idx = 1
        for line in text.split("\n"):
            if not line.strip():
                continue
            columns = split_into_columns(line)
            for col_idx, value in enumerate(columns, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)
            row_idx += 1

        wb.save(out_path)
    except pytesseract.TesseractNotFoundError:
        cleanup(src_path)
        return jsonify({
            "error": "Server पर Tesseract OCR install नहीं है। README देखें।"
        }), 500
    except Exception as e:
        cleanup(src_path)
        return jsonify({"error": f"Conversion असफल: {e}"}), 500

    cleanup(src_path)
    return send_file(
        out_path, as_attachment=True, download_name="extracted.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ---------------------------------------------------------------------------
# Frontend pages — हर tool का अपना URL (SEO के लिए ज़रूरी)
# ---------------------------------------------------------------------------

@app.route("/")
def home():
    return render_template("index.html", tools=TOOLS, site_url=SITE_URL)


@app.route("/<slug>")
def tool_page(slug):
    if slug in TOOLS:
        data = TOOLS[slug]
        return render_template(
            "tool_page.html", slug=slug, tool=data, all_tools=TOOLS,
            target_kb=None, site_url=SITE_URL,
        )
    if slug in COMPRESS_SIZE_PAGES:
        page = COMPRESS_SIZE_PAGES[slug]
        base_tool = dict(TOOLS["pdf-compress"])
        base_tool.update({
            "title": page["title"],
            "meta_description": page["meta_description"],
            "h1": page["h1"],
            "intro": page["intro"],
        })
        return render_template(
            "tool_page.html", slug=slug, tool=base_tool, all_tools=TOOLS,
            target_kb=page["target_kb"], site_url=SITE_URL,
        )
    from flask import abort
    abort(404)


@app.route("/sitemap.xml")
def sitemap():
    urls = [SITE_URL + "/"]
    urls += [f"{SITE_URL}/{slug}" for slug in TOOLS]
    urls += [f"{SITE_URL}/{slug}" for slug in COMPRESS_SIZE_PAGES]
    xml_items = "".join(f"<url><loc>{u}</loc></url>" for u in urls)
    xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">'
        f"{xml_items}</urlset>"
    )
    return Response(xml, mimetype="application/xml")


@app.route("/robots.txt")
def robots():
    lines = [
        "User-agent: *",
        "Allow: /",
        f"Sitemap: {SITE_URL}/sitemap.xml",
    ]
    return Response("\n".join(lines), mimetype="text/plain")


@app.route("/api/health")
def health():
    return jsonify({"status": "ok"})


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=False)
